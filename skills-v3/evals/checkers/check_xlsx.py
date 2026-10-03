#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""check_xlsx.py — kiểm tra CẤU TRÚC file Excel output của skill theo spec JSON (v3.4.0).
Dùng cho eval tầng 2 (LLM-in-loop): sau khi skill sinh file, chạy checker này thay vì soi mắt.

Usage: python3 check_xlsx.py <file.xlsx> <expected.json>
expected.json:
{
  "required_sheets": ["Assumptions", "Summary"],        // tên sheet PHẢI tồn tại (khớp đúng)
  "required_sheets_contains": ["Detail"],               // tồn tại sheet CHỨA chuỗi này
  "min_columns": { "Summary": 6 },                      // số cột header tối thiểu (đếm row có nhiều ô nhất trong 5 row đầu)
  "must_have_formula": ["Summary"],                     // sheet phải có ≥1 cell công thức (bắt đầu "=")
  "forbid_strings": ["fpt-canteen", "F-Pay", "Lorem"],  // chuỗi CẤM xuất hiện (bắt leak giá trị mẫu)
  "min_data_rows": { "Detail Estimate": 5 }             // số dòng dữ liệu tối thiểu
}
Exit 0 = PASS, 1 = FAIL (in từng lỗi)."""
import json, sys

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


def check(xlsx_path: str, exp: dict) -> list:
    import openpyxl
    errs = []
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    names = wb.sheetnames
    for s in exp.get("required_sheets", []):
        if s not in names:
            errs.append(f"thiếu sheet '{s}' (có: {names})")
    for sub in exp.get("required_sheets_contains", []):
        if not any(sub.lower() in n.lower() for n in names):
            errs.append(f"không có sheet nào chứa '{sub}' (có: {names})")
    for sheet, mincol in (exp.get("min_columns") or {}).items():
        tgt = next((n for n in names if n == sheet or sheet.lower() in n.lower()), None)
        if not tgt:
            continue
        ws = wb[tgt]
        width = max((sum(1 for c in row if c is not None) for row in
                     ws.iter_rows(min_row=1, max_row=5, values_only=True)), default=0)
        if width < mincol:
            errs.append(f"sheet '{tgt}': {width} cột < {mincol} yêu cầu")
    for sheet in exp.get("must_have_formula", []):
        tgt = next((n for n in names if n == sheet or sheet.lower() in n.lower()), None)
        if not tgt:
            errs.append(f"must_have_formula: không tìm thấy sheet '{sheet}'")
            continue
        has = any(isinstance(c.value, str) and c.value.startswith("=")
                  for row in wb[tgt].iter_rows() for c in row)
        if not has:
            errs.append(f"sheet '{tgt}': KHÔNG có công thức nào (tổng bị hardcode? — vi phạm excel-style.md)")
    forbid = exp.get("forbid_strings", [])
    if forbid:
        for n in names:
            for row in wb[n].iter_rows(values_only=True):
                for c in row:
                    if isinstance(c, str):
                        for f in forbid:
                            if f.lower() in c.lower():
                                errs.append(f"sheet '{n}': chứa chuỗi CẤM '{f}' trong {c[:60]!r} (leak giá trị mẫu?)")
    for sheet, minrows in (exp.get("min_data_rows") or {}).items():
        tgt = next((n for n in names if n == sheet or sheet.lower() in n.lower()), None)
        if not tgt:
            continue
        n_data = sum(1 for row in wb[tgt].iter_rows(min_row=2, values_only=True)
                     if any(c is not None for c in row))
        if n_data < minrows:
            errs.append(f"sheet '{tgt}': {n_data} dòng dữ liệu < {minrows} yêu cầu")
    return errs


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(2)
    errs = check(sys.argv[1], json.loads(open(sys.argv[2], encoding="utf-8-sig").read()))
    if errs:
        print(f"❌ {sys.argv[1]}: {len(errs)} lỗi cấu trúc:")
        for e in errs[:30]:
            print("   -", e)
        sys.exit(1)
    print(f"✅ {sys.argv[1]}: PASS cấu trúc")
    sys.exit(0)
