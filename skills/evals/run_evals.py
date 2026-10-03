#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""run_evals.py — eval harness TẦNG 1 (tự động, không cần LLM) cho bộ skill (v3.5.0).
Kiểm toàn bộ lớp deterministic: scripts dùng chung + gen Postman/OpenAPI + linter.
Chạy sau MỖI lần sửa script/skill, cùng với skill_doctor.

Usage: python3 run_evals.py [<skills_dir>]
Exit 0 = tất cả PASS, 1 = có FAIL. Tầng 2 (golden LLM cases): xem evals/README.md."""
import json, os, shutil, subprocess, sys, tempfile, time
from pathlib import Path

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

SKILLS = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else Path(__file__).resolve().parents[1]
PY = sys.executable
results = []  # (name, ok, msg)


def run(cmd, **kw):
    kw.setdefault("encoding", "utf-8")
    kw.setdefault("errors", "replace")
    return subprocess.run([PY] + cmd, capture_output=True, text=True, timeout=120, **kw)


def t(name):
    def deco(fn):
        def wrap():
            try:
                fn()
                results.append((name, True, ""))
            except AssertionError as e:
                results.append((name, False, str(e)))
            except Exception as e:
                results.append((name, False, f"{type(e).__name__}: {e}"))
        wrap.__name__ = name
        return wrap
    return deco


def make_ws(known_issue=False):
    ws = Path(tempfile.mkdtemp(prefix="eval_ws_"))
    ctx = {"version": "1.0",
           "project": {"name": "Eval Demo", "type": "web_app", "current_sprint": 1},
           "tech_stack": {"backend": "FastAPI"},
           "team": [{"id": "BE-1", "role": "Backend Dev", "allocation": 1.0}],
           "settings": {"working_days_per_week": 5},
           "known_issues": ([{"severity": "Critical", "status": "Open", "title": "x"}] if known_issue else []),
           "change_requests": [], "activity_log": []}
    (ws / "project-context.json").write_text(json.dumps(ctx, ensure_ascii=False, indent=2), encoding="utf-8")
    return ws


def make_qa_tracker(ws: Path, n_open_p1: int):
    import openpyxl
    wb = openpyxl.Workbook()
    s = wb.active
    s.append(["#", "Chủ đề", "Câu hỏi", "Priority", "Lý do", "Câu trả lời", "Người", "Ngày", "Trạng thái"])
    for i in range(n_open_p1):
        s.append([f"A{i}", "Auth", f"Q{i}?", "🔴 P1", "-", "", "", "", "⬜ Pending"])
    s.append(["B1", "Auth", "done?", "🔴 P1", "-", "OK", "PO", "01/07", "✅ Answered"])
    s.append(["C1", "NFR", "later?", "🟠 P2", "-", "", "", "", "⬜ Pending"])
    wb.save(ws / "Eval_Demo_QA_Tracker.xlsx")


@t("T1 update_context: set + append + atomic + validate")
def t1():
    ws = make_ws()
    r = run([str(SKILLS / "a1-project-init/scripts/update_context.py"), str(ws), "set", "estimates.total_md", "329"])
    assert r.returncode == 0, r.stdout + r.stderr
    r = run([str(SKILLS / "a1-project-init/scripts/update_context.py"), str(ws), "append", "activity_log",
             '{"skill":"estimate","date":"2026-07-09","outputs":["a.xlsx"]}'])
    assert r.returncode == 0, r.stdout + r.stderr
    ctx = json.loads((ws / "project-context.json").read_text(encoding="utf-8"))
    assert ctx["estimates"]["total_md"] == 329, "set không ăn"
    assert len(ctx["activity_log"]) == 1 and ctx["activity_log"][0]["skill"] == "estimate", "append không ăn"
    assert "updated_at" in ctx, "thiếu updated_at"
    assert not (ws / "project-context.json.lock").exists(), "lock không được dọn"
    shutil.rmtree(ws)


@t("T1b rotation: activity_log > 200 → archive JSONL, giữ 100")
def t1b():
    ws = make_ws()
    ctx = json.loads((ws / "project-context.json").read_text(encoding="utf-8"))
    ctx["activity_log"] = [{"skill": f"s{i}", "date": "2026-07-09"} for i in range(205)]
    (ws / "project-context.json").write_text(json.dumps(ctx), encoding="utf-8")
    r = run([str(SKILLS / "a1-project-init/scripts/update_context.py"), str(ws), "append", "activity_log",
             '{"skill":"last","date":"2026-07-09"}'])
    assert r.returncode == 0 and "rotate" in r.stdout, r.stdout
    ctx = json.loads((ws / "project-context.json").read_text(encoding="utf-8"))
    assert len(ctx["activity_log"]) == 100, f"log còn {len(ctx['activity_log'])}, phải 100"
    assert ctx["activity_log"][-1]["skill"] == "last", "entry mới nhất phải được giữ"
    arch = (ws / "project-context.archive.jsonl").read_text(encoding="utf-8").strip().splitlines()
    assert len(arch) == 106 and json.loads(arch[0])["skill"] == "s0", f"archive {len(arch)} dòng"
    shutil.rmtree(ws)


@t("T1c snapshot + restore: ghi nhầm → khôi phục được")
def t1c():
    ws = make_ws()
    r = run([str(SKILLS / "a1-project-init/scripts/update_context.py"), str(ws), "set", "project.name", "SAI"])
    assert r.returncode == 0, r.stdout
    assert list((ws / ".context-history").glob("context-*.json")), "không có snapshot"
    r = run([str(SKILLS / "a1-project-init/scripts/update_context.py"), str(ws), "restore"])
    assert r.returncode == 0, r.stdout + r.stderr
    ctx = json.loads((ws / "project-context.json").read_text(encoding="utf-8"))
    assert ctx["project"]["name"] == "Eval Demo", f"restore sai: {ctx['project']['name']}"
    shutil.rmtree(ws)


@t("T8 match_intent: map task → skill đúng + không đoán bừa")
def t8():
    r = run([str(SKILLS / "menu/scripts/match_intent.py"), str(SKILLS), "khách muốn đổi scope giữa sprint"])
    assert r.returncode == 0 and "/c7-change-request" in r.stdout.splitlines()[1], r.stdout
    r = run([str(SKILLS / "menu/scripts/match_intent.py"), str(SKILLS), "nấu phở bò"])
    assert r.returncode == 1 and "Không skill nào khớp" in r.stdout, "phải từ chối task ngoài phạm vi"


@t("T2a check_gate: 5 P1 mở → CHẶN (exit 1)")
def t2a():
    ws = make_ws()
    make_qa_tracker(ws, 5)
    r = run([str(SKILLS / "a1-project-init/scripts/check_gate.py"), str(ws)])
    assert r.returncode == 1, f"exit={r.returncode}, phải 1\n{r.stdout}"
    assert "GATE P1" in r.stdout and "5" in r.stdout, r.stdout
    shutil.rmtree(ws)


@t("T2b check_gate: 2 P1 mở + 0 bug → PASS (exit 0)")
def t2b():
    ws = make_ws()
    make_qa_tracker(ws, 2)
    r = run([str(SKILLS / "a1-project-init/scripts/check_gate.py"), str(ws)])
    assert r.returncode == 0, f"exit={r.returncode}\n{r.stdout}"
    shutil.rmtree(ws)


@t("T2c check_gate: bug Critical mở → CHẶN go-live")
def t2c():
    ws = make_ws(known_issue=True)
    r = run([str(SKILLS / "a1-project-init/scripts/check_gate.py"), str(ws)])
    assert r.returncode == 1 and "GATE BUG" in r.stdout, r.stdout
    shutil.rmtree(ws)


@t("T3 compute_status: staleness input mới hơn output")
def t3():
    ws = make_ws()
    now = time.time()
    for fname, age in [("Eval_Demo_Estimate.xlsx", 3600), ("Eval_Demo_API_Design.xlsx", 60),
                       ("Eval_Demo_DB_Design.xlsx", 60)]:
        f = ws / fname
        f.touch()
        os.utime(f, (now - age, now - age))
    r = run([str(SKILLS / "a1-project-init/scripts/compute_status.py"), str(ws), str(SKILLS)])
    assert r.returncode == 0, r.stderr
    d = json.loads(r.stdout)
    import re
    st = {}
    for x in d["documents"]:
        st[x["skill"]] = x["status"]
        st[re.sub(r'^[a-z]\d+-', '', x["skill"])] = x["status"]
    assert st.get("estimate") == "stale", f"estimate phải stale, được: {st.get('estimate')}"
    assert st.get("api-design") == "present", st.get("api-design")
    shutil.rmtree(ws)


@t("T4 check_circular_fk: bắt chu trình + pass đồ thị sạch")
def t4():
    p = Path(tempfile.mktemp(suffix=".json"))
    p.write_text('{"orders":["users"],"users":["orders"]}', encoding="utf-8")
    r = run([str(SKILLS / "a5-db-design/scripts/check_circular_fk.py"), str(p)])
    assert r.returncode == 1 and "CIRCULAR" in r.stdout, r.stdout
    p.write_text('{"orders":["users"],"users":[]}', encoding="utf-8")
    r = run([str(SKILLS / "a5-db-design/scripts/check_circular_fk.py"), str(p)])
    assert r.returncode == 0, r.stdout
    p.unlink()


@t("T5 gen_postman + gen_openapi: golden spec → cấu trúc đúng")
def t5():
    spec = SKILLS / "evals/golden/a4-api-design/api_spec.json"
    if not spec.exists():
        spec = SKILLS / "evals/golden/api-design/api_spec.json"
    tmp = Path(tempfile.mkdtemp(prefix="eval_gen_"))
    r = run([str(SKILLS / "a4-api-design/scripts/gen_postman.py"), str(spec), str(tmp / "c.json")])
    assert r.returncode == 0, r.stdout + r.stderr
    coll = json.loads((tmp / "c.json").read_text(encoding="utf-8"))
    folders = [f["name"] for f in coll["item"]]
    assert "⚙️ Setup & Auth" in folders and "Rooms" in folders, folders
    rooms_get = coll["item"][folders.index("Rooms")]["item"][0]["request"]
    assert any(h["key"] == "Authorization" for h in rooms_get["header"]), "thiếu Bearer header cho auth_required"
    r = run([str(SKILLS / "a4-api-design/scripts/gen_openapi.py"), str(spec), str(tmp / "o.yaml")])
    assert r.returncode == 0, r.stdout + r.stderr
    try:
        import yaml
        doc = yaml.safe_load((tmp / "o.yaml").read_text(encoding="utf-8"))
        assert doc["openapi"] == "3.0.3" and len(doc["paths"]) == 3, "openapi paths sai"
        assert "parameters" in doc["paths"]["/rooms/{roomId}/bookings"]["post"], "thiếu path param"
    except ImportError:
        pass
    shutil.rmtree(tmp)


@t("T6 skill_doctor: bộ skill hiện tại phải 0 FAIL")
def t6():
    r = run([str(SKILLS / "x6-skill-doctor/scripts/skill_doctor.py"), str(SKILLS)])
    assert r.returncode == 0, "skill_doctor có FAIL:\n" + r.stdout


@t("T7 check_xlsx: tự kiểm checker (bắt lỗi thật, pass file đúng)")
def t7():
    import openpyxl
    tmp = Path(tempfile.mkdtemp(prefix="eval_chk_"))
    wb = openpyxl.Workbook()
    s = wb.active; s.title = "Summary"
    s.append(["Category", "BD", "DD", "Dev", "UT", "IT", "Total"])
    s.append(["Backend", 1, 2, 10, 3, 2, "=SUM(B2:F2)"])
    wb.create_sheet("Detail Estimate").append(["No", "Task"])
    wb.save(tmp / "ok.xlsx")
    exp = tmp / "exp.json"
    exp.write_text(json.dumps({"required_sheets": ["Summary"], "must_have_formula": ["Summary"],
                               "min_columns": {"Summary": 7}, "forbid_strings": ["fpt-canteen"]}), encoding="utf-8")
    r = run([str(SKILLS / "evals/checkers/check_xlsx.py"), str(tmp / "ok.xlsx"), str(exp)])
    assert r.returncode == 0, r.stdout
    wb2 = openpyxl.Workbook()
    s2 = wb2.active; s2.title = "Summary"
    s2.append(["Category", "Total"]); s2.append(["Backend", 18])
    s2.append(["Ghi chú", "deploy tại fpt-canteen.vn"])
    wb2.save(tmp / "bad.xlsx")
    r = run([str(SKILLS / "evals/checkers/check_xlsx.py"), str(tmp / "bad.xlsx"), str(exp)])
    assert r.returncode == 1, "checker phải FAIL file thiếu công thức + leak chuỗi cấm"
    assert "công thức" in r.stdout and "fpt-canteen" in r.stdout, r.stdout
    shutil.rmtree(tmp)


if __name__ == "__main__":
    for fn in [t1, t1b, t1c, t2a, t2b, t2c, t3, t4, t5, t6, t7, t8]:
        fn()
    n_fail = sum(1 for _, ok, _ in results if not ok)
    print(f"\n🧪 EVALS TẦNG 1 — {len(results)} test | {len(results) - n_fail} PASS | {n_fail} FAIL\n")
    for name, ok, msg in results:
        print(("✅" if ok else "❌"), name, ("" if ok else f"\n     → {msg}"))
    if n_fail == 0:
        print("\n✅ Tầng 1 sạch. Tầng 2 (golden LLM cases): evals/README.md")
    sys.exit(1 if n_fail else 0)
